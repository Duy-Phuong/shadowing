"""Build ../vocabulary.html — the whole "total" sheet of Vocabulary-new.xlsx as one filterable, editable table.

Columns are kept as in Excel: ID | Name | Type | Description | Unit (Unit = level 0–6 / *).
Rows with no Name, Type, Description and Unit (placeholder IDs, blank rows) are skipped.
The page shell lives in vocabulary.template.html; the rows are embedded as JSON.
Usage: python3 build_total.py
"""
import json
from datetime import date, datetime
from pathlib import Path
import openpyxl

HERE = Path(__file__).parent
XLSX = HERE.parent / 'Vocabulary-new.xlsx'
TEMPLATE = HERE / 'vocabulary.template.html'
OUT = HERE.parent / 'vocabulary.html'

S = lambda c: '' if c is None else str(c).strip()


def main():
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)['total']
    stamp = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[5]
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        rid, name, typ, desc, unit = r[1:6]
        if not any(S(v) for v in (name, typ, desc, unit)):
            continue
        rows.append([rid, '' if name is None else str(name).rstrip(), S(typ), S(desc), S(unit)])
    day = stamp.date().isoformat() if isinstance(stamp, (datetime, date)) else date.today().isoformat()
    data = json.dumps({'date': day, 'rows': rows}, ensure_ascii=False, separators=(',', ':'))
    page = TEMPLATE.read_text(encoding='utf-8').replace('/*__DATA__*/', data.replace('</', '<\\/'))
    OUT.write_text(page, encoding='utf-8')
    print(f'{len(rows)} rows → {OUT} ({OUT.stat().st_size // 1024} KB)')


if __name__ == '__main__':
    main()
